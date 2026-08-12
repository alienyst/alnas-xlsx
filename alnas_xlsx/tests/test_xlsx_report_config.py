from base64 import b64encode
from io import BytesIO
from zipfile import ZipFile

from odoo.tests.common import TransactionCase
from openpyxl import Workbook, load_workbook


class TestXlsxReportConfig(TransactionCase):
    def test_copy_model_onchange_domain_and_idempotent_actions(self):
        model = self.env["ir.model"]._get("res.partner")
        field = self.env["ir.model.fields"].search(
            [("model_id", "=", model.id), ("name", "=", "name")], limit=1
        )
        config = self.env["xlsx.report.config"].create(
            {
                "name": "Partner Report",
                "report_name": "partner_report",
                "model_id": model.id,
                "field_id": field.id,
                "report_xlsx_template": b64encode(b"xlsx"),
                "report_xlsx_template_filename": "partner.xlsx",
                "domain": "[('active', '=', True)]",
            }
        )

        copied = config.copy_data()[0]
        self.assertEqual(copied["name"], "Partner Report (copy)")
        self.assertEqual(copied["report_name"], "partner_report (copy)")

        onchange = self.env["xlsx.report.config"].new({"model_id": model.id})
        onchange._onchange_model_id()
        self.assertEqual(onchange.field_id.model_id, model)
        self.assertEqual(onchange.field_id.ttype, "char")

        config._action_publish()
        config._action_publish()
        self.assertEqual(config.action_report_id.domain, config.domain)
        config._action_unpublish()
        config._action_unpublish()

    def test_multi_record_render_isolated_and_zip_names_safe(self):
        template = BytesIO()
        workbook = Workbook()
        workbook.active["A1"] = "{{ docs.name }}"
        workbook.save(template)

        report = self.env["ir.actions.report"].new(
            {
                "name": "Partner XLSX",
                "model": "res.partner",
                "report_type": "xlsx-jinja",
                "report_xlsx_jinja_template": b64encode(template.getvalue()),
                "report_xlsx_jinja_template_name": "partner.xlsx",
                "print_report_name": "'same'",
            }
        )
        partners = self.env["res.partner"].create(
            [{"name": "First/Partner"}, {"name": "Second\\Partner"}]
        )

        content, file_format = report._render_xlsx_records(partners.ids)

        self.assertEqual(file_format, "zip")
        with ZipFile(BytesIO(content)) as archive:
            self.assertEqual(archive.namelist(), ["same.xlsx", "same_2.xlsx"])
            values = [
                load_workbook(BytesIO(archive.read(name))).active["A1"].value
                for name in archive.namelist()
            ]
        self.assertEqual(values, partners.mapped("name"))
